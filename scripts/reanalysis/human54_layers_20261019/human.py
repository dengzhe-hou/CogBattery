"""Frozen human54 human references, CPU only; no model records are read.

All output participants are anonymous rows. NB uses stimulus episode IDs solely
for the prespecified paired-block bootstrap. Original source files stay unchanged.
"""
from __future__ import annotations
import os
for _key in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
    os.environ[_key] = '1'
import argparse
from collections import Counter
import json
from pathlib import Path
import re
import sys
import time
from datetime import datetime, timezone
import warnings
import numpy as np
import pandas as pd
from scipy.stats import rankdata
ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from scripts.reanalysis.sro_structure_20261015 import build_human as SRO
from scripts.reanalysis.itemfit_20261018 import strachan as IF_FB
from scripts.experiments.same_stimulus_20261011 import scoring as P9
from scripts.experiments.human13_20261017.memory import _eerland_materials
RAW = Path('data/raw')
OUT = ROOT / 'results/reanalysis/human54_layers_20261019/human_side'
DS_WEIGHTS = {'forward|3': 3, 'forward|4': 4, 'forward|5': 4, 'forward|6': 6, 'forward|7': 11,
              'backward|3': 3, 'backward|4': 5, 'backward|5': 8, 'backward|6': 12, 'backward|7': 8}
RULES = ('valid_only', 'invalid_wrong')


def clean(x):
    if isinstance(x, dict): return {str(k): clean(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)): return [clean(v) for v in x]
    if isinstance(x, np.ndarray): return clean(x.tolist())
    if isinstance(x, (float, np.floating)): return float(x) if np.isfinite(x) else None
    if isinstance(x, (int, np.integer)): return int(x)
    if isinstance(x, (bool, np.bool_)): return bool(x)
    return x


def auc(correct, confidence):
    y, x = np.asarray(correct, int), np.asarray(confidence, float)
    n1, n0 = int(y.sum()), int(len(y)-y.sum())
    return float((rankdata(x)[y == 1].sum()-n1*(n1+1)/2)/(n1*n0)) if n1 and n0 else np.nan


def fixed_weight(table, weights):
    """Require all frozen support; never renormalize around missing cells."""
    return SRO.weighted(table, weights, True)


class Reference:
    def __init__(self):
        self.endpoints, self.profiles, self.matrices, self.sources, self.notes = [], [], [], [], []

    def endpoint(self, task, endpoint, values=None, *, value=None, n=None, rule='both', cohort='primary', **extra):
        arr = np.asarray([] if values is None else values, float)
        arr = arr[np.isfinite(arr)]
        self.endpoints.append(dict(task=task, endpoint=endpoint, rule=rule, cohort=cohort,
            value=float(np.mean(arr)) if value is None and len(arr) else value,
            n=len(arr) if n is None else int(n), values=arr.tolist(), **extra))

    def profile(self, task, profile, cells, numerators, denominators, *, strata=None, ids=None,
                rule='both', cohort='primary', **extra):
        num, den = np.asarray(numerators, float), np.asarray(denominators, float)
        assert num.shape == den.shape and num.ndim == 2 and num.shape[1] == len(cells)
        assert np.isfinite(num).all() and np.isfinite(den).all()
        assert (den >= 0).all()
        value = np.divide(num.sum(0), den.sum(0), out=np.full(len(cells), np.nan), where=den.sum(0)>0)
        result = dict(task=task, profile=profile, rule=rule, cohort=cohort, cells=list(cells), values=value,
            unit_numerators=num, unit_denominators=den,
            unit_strata=list(strata) if strata is not None else ['all']*len(num), **extra)
        if ids is not None: result['unit_ids'] = list(ids)
        self.profiles.append(result)

    def mean_profile(self, task, name, table, **kwargs):
        x = table.to_numpy(float)
        self.profile(task, name, list(table.columns), np.nan_to_num(x), np.isfinite(x).astype(float),
                     aggregation='equal-person cell means', **kwargs)

    def matrix(self, name, table, *, cohort, rule, **extra):
        complete = table.dropna()
        self.matrices.append(dict(name=name, cohort=cohort, rule=rule, columns=list(table.columns),
            values=complete.to_numpy(float), n=len(complete), n_source=len(table),
            n_available_by_task=table.notna().sum().to_dict(), **extra))


def build_sro(ref):
    screened, qc = {}, {}
    for task in SRO.TASKS:
        d = SRO.load(task)
        flags = SRO.qc(task, d)
        qc[task] = dict(n_source=len(flags), n_pass=int(flags.passed.sum()))
        screened[task] = d[d.worker_id.isin(flags.index[flags.passed])].copy()
    ref.sources.append(dict(source=str(SRO.RAW), qc=qc, task_qc_implementation=str(Path(SRO.__file__))))
    for rule in RULES:
        matrices = pd.DataFrame(index=pd.Index(sorted(set().union(*(set(d.worker_id) for d in screened.values())))))
        ds = screened['digit_span'].copy()
        ds['direction'] = ds.condition.replace({'reverse': 'backward'})
        ds['cell'] = ds.direction + '|' + ds.num_digits.astype(int).astype(str)
        if rule == 'valid_only': ds = ds[ds.rt != -1]
        dc = ds.groupby(['worker_id','cell']).correct.mean().unstack('cell')
        cells = [f'{direction}|{length}' for direction in ('forward','backward') for length in range(3,10)]
        dc = dc.reindex(columns=cells)
        ref.mean_profile('digit_span','direction_length',dc,rule=rule)
        for cell in cells:
            direction, length = cell.split('|')
            ref.endpoint('digit_span',f'{direction}_length_{length}',dc[cell],rule=rule)
        # Contrasts match trial-weighted model summaries; each person contributes one contrast.
        dg = ds.groupby(['worker_id','direction']).correct.mean().unstack('direction')
        ref.endpoint('digit_span','forward_minus_backward',dg.forward-dg.backward,rule=rule)
        ds['length_band'] = np.where(ds.num_digits.isin([3,4]),'short',np.where(ds.num_digits.isin([5,6,7]),'long','excluded'))
        db = ds.groupby(['worker_id','length_band']).correct.mean().unstack('length_band')
        ref.endpoint('digit_span','short_minus_long',db['short']-db['long'],rule=rule)
        matrices['digit_span'] = fixed_weight(dc,DS_WEIGHTS)
        nb = screened['adaptive_n_back']
        nb = nb[(nb.exp_stage == 'adaptive') & nb.load.isin([2,3])].copy()
        if rule == 'valid_only': nb = nb[nb.rt != -1]
        blocks = nb.groupby(['worker_id','load','block_num']).correct.mean()
        nc = blocks.groupby(['worker_id','load']).mean().unstack('load')
        matrices['n_back'] = fixed_weight(nc,{2:1,3:1})
        for source, task, condition, weights in [
            ('stroop','stroop','condition',{'congruent':1,'incongruent':1}),
            ('attention_network_task','flanker','flanker_type',{'congruent':1,'incongruent':1,'neutral':1}),
            ('go_nogo','go_nogo','condition',{'go':945,'nogo':105})]:
            d = screened[source]; d = d[d.exp_stage != 'practice'].copy()
            if rule == 'valid_only' and task != 'go_nogo': d = d[d.rt != -1]
            dc = d.groupby(['worker_id',condition]).correct.mean().unstack(condition)
            matrices[task] = fixed_weight(dc,weights)
            if task != 'go_nogo':
                for c in weights: ref.endpoint(task,c+'_accuracy',dc[c],rule=rule)
                ref.endpoint(task,'error_interference',dc.congruent-dc.incongruent,rule=rule)
            if task == 'stroop':
                # Mapping is source-backed: response codes consistently map to ink colour.
                mapping = d.groupby('correct_response').stim_color.agg(lambda x: set(x))
                assert all(len(x)==1 for x in mapping)
                key_color = {k:next(iter(v)) for k,v in mapping.items()}
                inc = d[d.condition=='incongruent'].copy()
                inc['word_response'] = inc.key_press.map(key_color).str.lower() == inc.stim_word.str.lower()
                ref.endpoint(task,'word_response_on_incongruent',inc.groupby('worker_id').word_response.mean(),rule=rule)
            if task == 'flanker':
                d['cell'] = d.cue + '|' + d.flanker_type
                cells = [f'{cue}|{c}' for cue in ('center','double','nocue','spatial') for c in weights]
                ref.mean_profile(task,'cue_condition',d.groupby(['worker_id','cell']).correct.mean().unstack('cell').reindex(columns=cells),rule=rule)
                fl_no_neutral = fixed_weight(dc,{'congruent':1,'incongruent':1})
            if task == 'go_nogo':
                vals = []
                for _, g in d.groupby('worker_id'):
                    go, press = g.condition.eq('go'),g.key_press.eq(32)
                    h, ns, f, nn = int((go & press).sum()),int(go.sum()),int((~go & press).sum()),int((~go).sum())
                    vals.append(dict(hit_rate=h/ns,false_alarm_rate=f/nn,omission_rate=1-h/ns,
                                     dprime=SRO.dprime(h,ns,f,nn),balanced_accuracy=(h/ns+1-f/nn)/2))
                vals = pd.DataFrame(vals)
                for c in vals: ref.endpoint(task,c,vals[c],rule=rule)
        ref.matrix('sro',matrices,cohort='sro',rule=rule,structure='sro',representation='full',neutral='included',weights=dict(digit_span=DS_WEIGHTS,n_back={'2':1,'3':1},stroop={'congruent':1,'incongruent':1},flanker={'congruent':1,'incongruent':1,'neutral':1},go_nogo={'go':945,'nogo':105}))
        matrices_no = matrices.copy(); matrices_no['flanker'] = fl_no_neutral
        ref.matrix('sro_no_neutral',matrices_no,cohort='sro',rule=rule,structure='sro',representation='no_neutral',neutral='excluded')


def build_nback(ref):
    base = ROOT/'results/same_stimulus_20261011/materials'
    episodes = [json.loads(x) for x in (base/'nback_episodes.jsonl').read_text().splitlines() if x.strip()]
    humans = {x['episode_id']:x for x in map(json.loads,(base/'human_nback_trials.jsonl').read_text().splitlines())}
    assert len(episodes)==80
    for rule in ('valid_only','invalid_wrong','nb_r2'):
        cells = [f'{load}|{kind}' for load in (2,3) for kind in ('target','lure','control')]
        nums, dens, ids, strata = [],[],[],[]
        all_blocks = {2:[],3:[]}
        for ep in episodes:
            h = humans[ep['episode_id']]
            tr = []
            for kind,code in zip(ep['trial_kinds'],h['codes']):
                if code is None and rule=='invalid_wrong': code = kind != 'target'
                elif code is None and rule=='nb_r2': code = False
                tr.append(dict(kind=kind,code=code))
            all_blocks[ep['load']].append(tr)
            n, d = np.zeros(6),np.zeros(6)
            for j,cell in enumerate(cells):
                load,kind = cell.split('|')
                if ep['load']==int(load):
                    r=P9._rate(tr,kind,'R1');n[j]=r['k'];d[j]=r['n']
            nums.append(n);dens.append(d);ids.append(ep['episode_id']);strata.append(str(ep['load']))
        ref.profile('n_back','load_kind',cells,nums,dens,strata=strata,ids=ids,rule=rule,
                    aggregation='pooled trials; paired whole-block bootstrap within load')
        for load, blocks in all_blocks.items():
            ep = P9.endpoints(blocks)
            for kind in ('target','lure','control','nontarget'):
                values = [P9._rate(b,kind,'R1')['rate'] for b in blocks]
                ref.endpoint('n_back',f'load_{load}_{kind}',[np.nan if x is None else x for x in values],value=ep['R1'][kind]['rate'],rule=rule,n_trials=ep['R1'][kind]['n'],aggregation='pooled trials')
            dps=[P9.episode_dprime(b) for b in blocks]
            ref.endpoint('n_back',f'load_{load}_dprime_episode_mean',[np.nan if x is None else x for x in dps],rule=rule)
    ref.sources.append(dict(task='n_back',source=str(base/'human_nback_trials.jsonl'),blocks=80,trials=1800))


def build_ball(ref):
    path=ROOT/'results/human13_20261017/prepared_54_drm_sentence/ball_sentence_reference/ball_sentence_scores.csv'
    d=pd.read_csv(path)
    ref.endpoint('operation_span','position_accuracy',d.operation_span)
    names={'hit_rate':'drm_sentence_hit_rate','lure_false_alarm_rate':'drm_sentence_lure_fa','new_false_alarm_rate':'drm_sentence_unrelated_fa','delta_dprime':'drm_sentence_delta_dprime','native':'drm_sentence_native','accuracy_analyzed56':'drm_sentence_accuracy_analyzed56','balanced_accuracy':'drm_sentence_balanced_accuracy'}
    cells=[f'{w}|{k}' for w in ('no_warning','warning') for k in ('old','CL','newunrelated')]
    num,den=[],[]
    drm=d[d.drm_sentence_native.notna()].copy()
    for _,row in drm.iterrows():
        a,b=np.zeros(6),np.zeros(6)
        for j,(cell,col) in enumerate(zip(cells,[names[k] for _ in range(2) for k in ('hit_rate','lure_false_alarm_rate','new_false_alarm_rate')])):
            if cell.startswith(row.warning+'|') and pd.notna(row[col]):a[j]=row[col];b[j]=1
        num.append(a);den.append(b)
    ref.profile('drm_false_memory','warning_probe',cells,num,den,strata=drm.warning.tolist(),aggregation='equal-person cell means; whole-person resampling within warning condition')
    for warning in ('no_warning','warning'):
        for endpoint,col in names.items():ref.endpoint('drm_false_memory',warning+'|'+endpoint,d.loc[d.warning==warning,col])
    for layer in ('all','no_warning','warning'):
        g=d if layer=='all' else d[d.warning==layer]
        for rep,col in [('native','drm_sentence_native'),('accuracy_analyzed56','drm_sentence_accuracy_analyzed56'),('balanced_accuracy','drm_sentence_balanced_accuracy')]:
            m=g[['operation_span',col,'source_monitoring']].rename(columns={col:'drm_false_memory'})
            ref.matrix('ball_'+layer+'_'+rep,m,cohort='ball',rule='both',structure='ball_'+('pooled' if layer=='all' else layer),layer=layer,representation={'native':'lure_contrast','accuracy_analyzed56':'accuracy56','balanced_accuracy':'balanced_accuracy'}[rep])
    ref.sources.append(dict(task='ball',source=str(path),n_rows=len(d),n_drm=int(d.drm_sentence_native.notna().sum())))


def build_cvlt(ref):
    base=RAW/'ucla_cnp_ds000030'
    d=pd.read_csv(base/'phenotype/cvlt.tsv',sep='\t').merge(pd.read_csv(base/'participants.tsv',sep='\t')[['participant_id','diagnosis']],on='participant_id',validate='one_to_one')
    mapping={**{f'learning_{i}':f'cvlt_cor{i}' for i in range(1,6)},'learning_total':'cvlt_totcor','interference_recall':'cvlt_bf','short_free':'cvlt_sdf','long_free':'cvlt_ldf','short_cued':'cvlt_sdc','long_cued':'cvlt_ldc','free_intrusions':'cvlt_fri','cued_intrusions':'cvlt_cri','repetitions':'cvlt_tr','learning_slope':'cvlt_ls'}
    learning=d[[f'cvlt_cor{i}' for i in range(1,6)]].to_numpy(float)
    slopes=(learning @ np.arange(-2,3))/10
    np.testing.assert_allclose(slopes,d.cvlt_ls,atol=1e-12,rtol=0)
    for cohort,g in [('primary',d[d.diagnosis=='CONTROL']),('all_diagnoses',d)]:
        for endpoint,col in mapping.items():ref.endpoint('cvlt_word_list',endpoint,g[col],cohort=cohort)
        ref.endpoint('cvlt_word_list','recognition_hit_rate',g.cvlt_ldh/16,cohort=cohort)
        ref.endpoint('cvlt_word_list','recognition_false_alarm_rate',g.cvlt_ldfp/32,cohort=cohort)
        ref.endpoint('cvlt_word_list','learning_5_minus_1',g.cvlt_cor5-g.cvlt_cor1,cohort=cohort)
    ref.sources.append(dict(task='cvlt_word_list',source=str(base/'phenotype/cvlt.tsv'),diagnosis_counts=d.diagnosis.value_counts().to_dict(),excluded_endpoint='cvlt_rd formula unverified, frozen exclusion'))


def parse_source_answer_keys(rows, question_map):
    """Read the current experiment's author formulas; numeric labels are not shared."""
    head, labels = rows[:2]
    gold = {}
    for j, text in enumerate(labels):
        if text not in question_map:
            continue
        score_column = head.index('C'+head[j])
        formula = next(row[score_column] for row in rows[2:]
                       if isinstance(row[score_column], str) and row[score_column].startswith('=IF('))
        code = int(re.search(r'=(\d+),1,0\)', formula).group(1))
        gold[question_map[text]['question']] = code
    return gold


def build_source(ref,out,experiments=(1,2),write_trials=True):
    import openpyxl
    raw=RAW/'eerland2018_source_memory'
    _,speakers,questions=_eerland_materials(raw)
    qmap={q['question']:q for q in questions}
    # The administered export fixes four wording/typographic variants in Stimuli_design.
    question_aliases={
        'Who hoped they would all make it to the pub in time?':'Who hoped they all made it to the pub in time?',
        'Who wanted to have a Belgian beer?':'Who wanted to have a Belgium beer?',
        'Who thought the flag was of Romania?':'Who thought the flag was of Romenia?',
        'Who was walking from the pub?':'Who was going to walk home?'}
    for administered,design in question_aliases.items(): qmap[administered]=qmap[design]
    variants=('indirect','direct_agent_last','direct_agent_first')
    cells=[f'{v}|{q["source_row"]}' for v in variants for q in questions]
    anonymous_trials=[]
    for exp in experiments:
        # Exp2 uses Jake=3,Tammy=2,Olivia=1,Mike=4; Exp1 uses 1,2,3,4.
        # Read each author's file instead of transferring response codes between experiments.
        analysis_path=raw/f'Data_AnalysisExp{exp}.xlsx'
        book=openpyxl.load_workbook(analysis_path,read_only=True,data_only=False)
        gold=parse_source_answer_keys(list(book['Data file'].values),qmap)
        book.close()
        assert len(gold)==48
        speaker_codes={q['answer']:gold[q['question']] for q in questions}
        assert len(speaker_codes)==4 and all(speaker_codes[q['answer']]==gold[q['question']] for q in questions)
        book=openpyxl.load_workbook(raw/f'RawData_Exp{exp}.xlsx',read_only=True,data_only=True)
        rows=list(book.worksheets[0].values);head,labels=rows[:2]
        qindices={j:qmap[text] for j,text in enumerate(labels) if text in qmap}
        assert len(qindices)==48
        timings=[('Q91_3','indirect'),('Q96_3','direct_agent_last'),('Q100_3','direct_agent_first')]
        numerators,denominators,strata,condition_scores=[],[],[],[]
        counts=Counter()
        for row in rows[2:]:
            answered=[(j,q,row[j]) for j,q in qindices.items() if row[j] is not None]
            counts[f'answered_{len(answered)}']+=1
            if len(answered)!=24:continue
            assert all(code in (1,2,3,4) for _,_,code in answered)
            condition=[variant for field,variant in timings if field in head and row[head.index(field)] is not None]
            assert len(condition)==1
            variant=condition[0]
            num,den=np.zeros(len(cells)),np.zeros(len(cells))
            for j,q,code in answered:
                ix=cells.index(f'{variant}|{q["source_row"]}')
                num[ix]=int(code==gold[q['question']]);den[ix]=1
                anonymous_trials.append(dict(experiment=exp,anonymous_record=len(numerators),condition=variant,question_source_row=q['source_row'],answer_code=code,correct=int(num[ix])))
            numerators.append(num);denominators.append(den);strata.append(variant);condition_scores.append((variant,num.sum()/24))
        book.close()
        cohort='primary' if exp==1 else 'exp2'
        for variant in variants:
            values=[x for v,x in condition_scores if v==variant]
            if values:ref.endpoint('source_monitoring',variant+'_accuracy',values,cohort=cohort)
        active=[j for j,cell in enumerate(cells) if cell.split('|')[0] in set(strata)]
        ref.profile('source_monitoring','item_condition',[cells[j] for j in active],np.asarray(numerators)[:,active],np.asarray(denominators)[:,active],strata=strata,cohort=cohort,
                    aggregation='pooled trials per item; whole-person resampling within quotation condition')
        ref.sources.append(dict(task='source_monitoring',experiment=exp,source=str(raw/f'RawData_Exp{exp}.xlsx'),complete24=len(numerators),conditions=dict(Counter(strata)),response_count_distribution=dict(counts),selection='exactly 24 answered source questions; no paper exclusion-chain reconstruction',speaker_code_verification='48 experiment-specific author scoring formulas agree with material answer keys',answer_coding_source=str(analysis_path),speaker_codes=speaker_codes,question_wording_aliases=question_aliases))
    trials=pd.DataFrame(anonymous_trials)
    if write_trials: trials.to_csv(out/'source_monitoring_anonymous_trials.csv',index=False)
    return trials


def build_fb(ref):
    spec=json.loads((ROOT/'scripts/experiments/same_stimulus_20261011/SPEC.json').read_text())
    source=Path(spec['materials']['strachan']['source_human_scores'])
    h=IF_FB.human_reference(IF_FB.read_csv(source))
    main=[x for x in h.items if x['main_analysis']]
    cells=[x['item_id'] for x in main]
    nums,dens,strata=[],[],[]
    state_vals={'F':[],'T':[],'catch':[]}
    main_values=[]
    for lst,group in h.lists.items():
        for row in group.scores:
            num,den=np.zeros(len(cells)),np.zeros(len(cells))
            by_state={'F':[],'T':[],'catch':[]}
            for item_id,score in zip(group.item_ids,row):
                meta=next(x for x in h.items if x['item_id']==item_id)
                by_state[meta['state'] if meta['state'] in ('F','T') else 'catch'].append(score)
                if item_id in cells:num[cells.index(item_id)]=score;den[cells.index(item_id)]=1
            for state,vals in by_state.items():state_vals[state].append(np.mean(vals))
            main_values.append(np.mean(by_state['F']+by_state['T']))
            nums.append(num);dens.append(den);strata.append(lst)
    # Point endpoint is pooled responses, as in P9, with individual values retained for percentiles.
    for condition in ('session','isolated','session_no_system','isolated_no_system'):
        for state,label in [('F','FB'),('T','TB'),('catch','catch')]:
            selected=[x for x in h.items if (x['state']==state if state!='catch' else not x['main_analysis'])]
            value=sum(x['human_correct'] for x in selected)/sum(x['human_n'] for x in selected)
            ref.endpoint('false_belief',f'{condition}|{label}_accuracy',state_vals[state],value=value,aggregation='pooled item responses')
        ref.endpoint('false_belief',condition+'|main_accuracy',main_values)
        fb_items=[x for x in h.items if x['state']=='F']
        tb_items=[x for x in h.items if x['state']=='T']
        fb_rate=sum(x['human_correct'] for x in fb_items)/sum(x['human_n'] for x in fb_items)
        tb_rate=sum(x['human_correct'] for x in tb_items)/sum(x['human_n'] for x in tb_items)
        ref.endpoint('false_belief',condition+'|FB_minus_TB',np.array(state_vals['F'])-np.array(state_vals['T']),value=fb_rate-tb_rate,aggregation='difference of pooled FB and TB accuracy; individual paired differences retained')
        ref.profile('false_belief',condition+'_items',cells,nums,dens,strata=strata,aggregation='pooled item responses; whole-person resampling within A/B list')
    ref.sources.append(dict(task='false_belief',source=str(source),participants=sum(len(g.scores) for g in h.lists.values()),lists={k:len(v.scores) for k,v in h.lists.items()}))


def build_epitome(ref):
    source=RAW/'epitome2024_sn7gj/recursive_mindreading/tom_rm_trial.csv'
    d=pd.read_csv(source);d=d[d.response_type=='2AFC'].copy()
    assert len(d)==3315 and d.participant_id.nunique()==63
    cells=[f'{t}|{l}' for t in ('mental','control') for l in range(1,8)]
    d['cell']=d.question_type+'|'+d.question_level.astype(str)
    num=d.groupby(['participant_id','cell']).accuracy.sum().unstack('cell').reindex(columns=cells).fillna(0)
    den=d.groupby(['participant_id','cell']).accuracy.size().unstack('cell').reindex(columns=cells).fillna(0)
    ref.profile('epitome_tom','type_level',cells,num,den,aggregation='pooled trial accuracy; whole-person resampling')
    for cell in cells:
        g=d[d.cell==cell]
        ref.endpoint('epitome_tom',cell+'_accuracy',g.groupby('participant_id').accuracy.mean(),value=g.accuracy.mean(),n_trials=len(g),aggregation='pooled trial accuracy')
    for field in ('question_format','story_format'):
        for level,g in d.groupby(field):
            ref.endpoint('epitome_tom',field+'|'+str(level)+'_accuracy',g.groupby('participant_id').accuracy.mean(),value=g.accuracy.mean(),n_trials=len(g))
    for typ in ('mental','control'):
        profile=d[d.question_type==typ].groupby('question_level').accuracy.mean().reindex(range(1,8))
        ref.endpoint('epitome_tom',typ+'_level_slope',value=float(np.polyfit(np.arange(1,8),profile,1)[0]),n=63)
    ref.sources.append(dict(task='epitome_tom',source=str(source),n_rows=len(d),n_people=63,additional_filter='none'))


def build_cal(ref):
    import openpyxl
    raw=RAW/'mazancieux2020_metacognition'
    d=pd.read_csv(raw/'Data_RCJ_4tasks.csv',sep=';',encoding='latin-1');d=d[d['Procedure[Trial]']=='MS'].copy()
    book=openpyxl.load_workbook(raw/'Material_EM_SM_EF_tasks.xlsx',read_only=True,data_only=True)
    material=list(book['SM task'].values)[1:];book.close()
    lookup={frozenset((str(target).strip(),str(distractor).strip())):(i,str(target).strip()) for i,(_,target,distractor) in enumerate(material,1)}
    mapped=[lookup[frozenset((str(row.Gauche).strip(),str(row.Droite).strip()))] for row in d.itertuples()]
    d['item_id']=[x[0] for x in mapped]
    d['conf']=pd.to_numeric(d['ConfMS.RESP'].replace({'c':'10'}))
    d['correct']=(d['QuestMS.RESP']==d.BR).astype(int)
    expected=['s' if str(left).strip()==target else 'l' for left,(_,target) in zip(d.Gauche,mapped)]
    assert np.all(np.array(expected)==d.BR.to_numpy())
    assert len(d)==7240 and d.Subject.nunique()==181 and d.groupby('Subject').item_id.nunique().eq(40).all()
    assert d.conf.between(0,10).all()
    values=[]
    for _,g in d.groupby('Subject'):
        values.append(dict(accuracy=g.correct.mean(),confidence_mean_percent=g.conf.mean()*10,
            official_bias=(.5+.05*g.conf).mean()-g.correct.mean(),type2_auc_fixed_direction=auc(g.correct,g.conf)))
    values=pd.DataFrame(values)
    for endpoint in values:ref.endpoint('confidence_calibration',endpoint,values[endpoint])
    for name,col,scale in [('item_accuracy','correct',1),('item_confidence','conf',10)]:
        table=d.pivot(index='Subject',columns='item_id',values=col).reindex(columns=range(1,41))*scale
        table.columns=table.columns.astype(str)
        ref.mean_profile('confidence_calibration',name,table)
    ref.sources.append(dict(task='confidence_calibration',source=str(raw/'Data_RCJ_4tasks.csv'),n_people=181,n_trials=7240,confidence='raw integer 0..10; percent = 10*k; official bias uses .5+.05*k',item_linkage='unordered Target/Distractor pairs, target side checked against every BR'))


def build_wag(ref):
    ref.endpoint('post_decision_wagering','accuracy',value=.65,n=40,se=.019)
    ref.endpoint('post_decision_wagering','high_minus_low_accuracy',value=.105,n=40,se=.0218)
    ref.sources.append(dict(task='post_decision_wagering',source=str(RAW/'postdecision_wagering_text_candidates/dienes_seth2010.pdf'),page=677,n=40,data='published group mean and SE only',excluded_endpoint='logistic Type-II d-prime formula unverified, frozen exclusion'))


def build_required_contrasts(ref):
    """Archive the frozen direction contrasts from their source-native means.

    Between-condition contrasts have no fictitious individual paired values.
    """
    lookup={(r['task'],r['rule'],r['cohort'],r['endpoint']):r for r in ref.endpoints}
    def contrast(task,name,positive,negative,rule='both',cohort='primary',n=None,note='difference of source endpoint means'):
        pos=[lookup[task,rule,cohort,k] for k in positive]
        neg=[lookup[task,rule,cohort,k] for k in negative]
        value=np.mean([r['value'] for r in pos])-np.mean([r['value'] for r in neg])
        ref.endpoint(task,name,value=float(value),n=n or max(r['n'] for r in pos+neg),rule=rule,cohort=cohort,contrast_definition=note)
    for rule in ('valid_only','invalid_wrong','nb_r2'):
        for load in (2,3):
            contrast('n_back',f'load_{load}_lure_minus_control',[f'load_{load}_lure'],[f'load_{load}_control'],rule=rule,n=40)
            contrast('n_back',f'load_{load}_hit_minus_lure',[f'load_{load}_target'],[f'load_{load}_lure'],rule=rule,n=40)
    contrast('drm_false_memory','warning_minus_no_warning_lure',['warning|lure_false_alarm_rate'],['no_warning|lure_false_alarm_rate'],n=437,note='between-group warning minus no-warning; model contrast is within model')
    contrast('source_monitoring','direct_minus_indirect',['direct_agent_first_accuracy','direct_agent_last_accuracy'],['indirect_accuracy'],n=299,note='equal mean of two direct quotation conditions minus indirect; between participants')
    contrast('source_monitoring','direct_agent_first_minus_indirect',['direct_agent_first_accuracy'],['indirect_accuracy'],cohort='exp2',n=253,note='supplement: direct-agent-first minus indirect, between participants; model direct contrast averages its two direct conditions')
    for level in range(1,8):
        contrast('epitome_tom',f'level_{level}_mental_minus_control',[f'mental|{level}_accuracy'],[f'control|{level}_accuracy'],n=63)


def build(out):
    start,cpu=time.monotonic(),time.process_time()
    out.mkdir(parents=True,exist_ok=True)
    ref=Reference()
    for fn in (build_sro,build_nback,build_ball,build_cvlt):
        fn(ref);print(fn.__name__+' complete',flush=True)
    build_source(ref,out);print('build_source complete',flush=True)
    for fn in (build_fb,build_epitome,build_cal,build_wag):
        fn(ref);print(fn.__name__+' complete',flush=True)
    build_required_contrasts(ref)
    resource=dict(wall_seconds=time.monotonic()-start,cpu_seconds=time.process_time()-cpu,**SRO.memory_usage(),threads={k:os.environ[k] for k in ('OMP_NUM_THREADS','MKL_NUM_THREADS','OPENBLAS_NUM_THREADS')})
    data=clean(dict(status='human_side_completed_before_any_formal_behavior_read',schema_version=1,completed_at_utc=datetime.now(timezone.utc).isoformat(),endpoints=ref.endpoints,profiles=ref.profiles,matrices=ref.matrices,sources=ref.sources,notes=ref.notes,resource=resource))
    (out/'references.json').write_text(json.dumps(data,indent=2,allow_nan=False)+'\n')
    pd.DataFrame([{k:v for k,v in r.items() if k!='values'} for r in data['endpoints']]).to_csv(out/'endpoints.csv',index=False)
    (out/'manifest.json').write_text(json.dumps(dict(status=data['status'],completed_at_utc=data['completed_at_utc'],resource=resource,sources=data['sources'],endpoints=len(ref.endpoints),profiles=len(ref.profiles),matrices=[{k:v for k,v in m.items() if k!='values'} for m in data['matrices']]),indent=2)+'\n')
    print(json.dumps(resource),flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--out',type=Path,default=OUT);args=parser.parse_args();build(args.out)
