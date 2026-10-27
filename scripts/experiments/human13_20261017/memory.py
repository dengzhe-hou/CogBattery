"""Public-original memory tasks. Preparation only; no model/network/job code.

CVLT uses an explicitly prepared, source-verified Standard-form JSON material file.
See CVLT_MATERIAL_SCHEMA below and cvlt_standard_source.json. Synthetic fixtures are tests,
not research material. DRM sentence preparation is independent of the unverified
CITY eight-word set; the original combined word/sentence protocol stays blocked.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import json
from pathlib import Path
import random
import re
from statistics import NormalDist

TASKS = ["cvlt_word_list", "drm_false_memory", "source_monitoring"]
SYSTEM = "Follow the task instructions and return only the requested answer."
HISTORY_MODE = "full"

CVLT_MATERIAL_SCHEMA = {
    "version": "CVLT-II", "form": "Standard",
    "source": {"reference": "Original material reference and page numbers",
               "material_path": "Path to verified original source",
               "use_authorization": "User authorization for this workflow; not a copyright license"},
    "lists": {"A": [{"word": "original word", "category": "original category"}],
              "B": [{"word": "original word", "category": "original category"}]},
    "cue_order": ["four original A categories in administration order"],
    "recognition": [{"word": "original probe", "is_old": True}],
    "prompts": {"study": "Authorized instruction with {word}",
                "recall_a": "Instruction for an A learning-trial free recall",
                "recall_b": "Instruction for B free recall",
                "short_free": "Short-interval free recall instruction",
                "long_free": "Long-stage free recall instruction",
                "cued": "Authorized category instruction with {category}",
                "recognition": "Authorized yes/no instruction with {word}"},
    "protocol_reference": "Verified administration rules and text adaptation notes",
    "gap_steps": {"before_long": [], "before_forced_choice": []},
    "extra_steps": [],
    "forced_choice": [],
}


def _step(sid, prompt, kind, gold=None, **meta):
    return {"id": sid, "prompt": prompt, "kind": kind, "gold": gold,
            "meta": meta, "max_tokens": 256 if kind in ("cvlt_recall", "cvlt_recall_followup") else 24}


def _episode(eid, task, condition, steps, **extra):
    return {"id": eid, "task": task, "condition": condition, "system": SYSTEM,
            "history_mode": HISTORY_MODE, "steps": steps, **extra}


def _result(episodes, provenance, blockers=()):
    return {"status": "blocked" if blockers else "ready", "blockers": list(blockers),
            "provenance": provenance, "episodes": episodes}


def _norm(text):
    return " ".join(str(text).strip().casefold().split())


def _answer_text(response):
    """Remove complete formatting wrappers, never extract an answer from prose."""
    text = str(response).strip()
    reply = re.fullmatch(r"Reply:[ \t]*([^\r\n]+)", text, re.I)
    if reply:
        text = reply.group(1).strip()
    for opening, closing in (("**", "**"), ('"', '"'), ("'", "'"), ("`", "`")):
        if len(text) > len(opening) + len(closing) and text.startswith(opening) and text.endswith(closing):
            return text[len(opening):-len(closing)].strip()
    return text


def _label_with_option(text, options):
    """Accept a complete label plus its own option text, including wrong choices."""
    for key, value in options.items():
        label, option = re.escape(str(key)), re.escape(str(value))
        if (re.fullmatch(label + r"[.)]\s*" + option + r"[.!]?", text, re.I)
                or re.fullmatch(option + r"\s*\((?:option\s*)?" + label + r"\)[.!]?", text, re.I)):
            return key
    return None


def _choice(response, options):
    """Exact label with harmless surrounding punctuation; no guessed prose answer."""
    text = _norm(response).strip(" .!\"'`\n\t")
    return options.get(text)


def _load_cvlt(path):
    data = json.loads(Path(path).read_text())
    if data.get("version") != "CVLT-II" or data.get("form") != "Standard":
        raise ValueError("CVLT materials must explicitly identify CVLT-II Standard")
    for key in ("reference", "material_path", "use_authorization"):
        if not str(data.get("source", {}).get(key, "")).strip():
            raise ValueError(f"CVLT source.{key} is required; presence is not independent license verification")
    if not data.get("protocol_reference"):
        raise ValueError("CVLT protocol_reference is required")
    for name in ("A", "B"):
        items = data.get("lists", {}).get(name, [])
        if len(items) != 16 or any(not x.get("word") or not x.get("category") for x in items):
            raise ValueError(f"CVLT List {name} needs 16 original word/category items")
        if len({_norm(x["word"]) for x in items}) != 16:
            raise ValueError(f"CVLT List {name} contains duplicate words")
        if sorted(Counter(x["category"] for x in items).values()) != [4, 4, 4, 4]:
            raise ValueError(f"CVLT List {name} needs four categories with four words each")
    if set(data.get("cue_order", [])) != {x["category"] for x in data["lists"]["A"]} or len(data["cue_order"]) != 4:
        raise ValueError("CVLT cue_order must list the four A categories once")
    probes = data.get("recognition", [])
    if len(probes) != 48 or any(type(x.get("is_old")) is not bool or not x.get("word") for x in probes):
        raise ValueError("CVLT recognition needs 48 original probes with Boolean is_old")
    if len({_norm(x["word"]) for x in probes}) != 48:
        raise ValueError("CVLT recognition probes must be distinct")
    if {_norm(x["word"]) for x in probes if x["is_old"]} != {_norm(x["word"]) for x in data["lists"]["A"]}:
        raise ValueError("CVLT recognition targets must be exactly the 16 A words")
    for key in CVLT_MATERIAL_SCHEMA["prompts"]:
        if not data.get("prompts", {}).get(key):
            raise ValueError(f"CVLT prompts.{key} is required")
    for key, field in (("study", "{word}"), ("cued", "{category}"), ("recognition", "{word}")):
        if field not in data["prompts"][key]:
            raise ValueError(f"CVLT prompts.{key} must contain {field}")
    forced = data.get("forced_choice", [])
    if forced and (len(forced) != 16 or any(len(x.get("choices", [])) != 2 or x.get("answer") not in (1, 2) for x in forced)):
        raise ValueError("Optional CVLT forced_choice needs 16 two-choice items, answer 1 or 2")
    for gap in ("before_long", "before_forced_choice"):
        for item in data.get("gap_steps", {}).get(gap, []):
            if not item.get("prompt"):
                raise ValueError(f"CVLT gap step {gap} requires an explicit prompt")
    for item in data.get("extra_steps", []):
        if not item.get("prompt") or item.get("before") not in {"A1", "A2", "A3", "A4", "A5", "B", "short", "long", "recognition", "forced_choice"}:
            raise ValueError("CVLT extra step needs prompt and a recognized before stage")
    return data


def build_cvlt(config):
    path = Path(config.get("cvlt_materials", Path(config["materials_dir"]) / "cvlt_standard.json"))
    provenance = {"cohort": "CNP ds000030", "human_form": "272/272 Standard adult CVLT-II",
                  "material_file": str(path), "human_scores": str(Path(config["raw_root"]) / "ucla_cnp_ds000030/phenotype/cvlt.tsv"),
                  "schema": CVLT_MATERIAL_SCHEMA,
                  "limitations": ["Text replaces auditory presentation; full prior context remains accessible.",
                                  "Long-stage label is not an equivalent human elapsed delay.",
                                  "Human word-level logs are not required for matched aggregate recall endpoints."]}
    if not path.is_file():
        provenance["verified_source"] = str(Path(__file__).with_name("cvlt_standard_source.json"))
        return _result([], provenance, ["Verified adult CVLT-II Standard source is available, but the prepared material file is absent at the requested path."])
    data = _load_cvlt(path)
    provenance.update({"source": data["source"], "protocol_reference": data["protocol_reference"],
                       "long_gap_outputs": len(data.get("gap_steps", {}).get("before_long", [])),
                       "forced_choice_included": bool(data.get("forced_choice")),
                       "adaptation": data.get("adaptation", {}),
                       "recall_followup_included": bool(data.get("recall_followup_prompt"))})
    steps = []
    words_a = [x["word"] for x in data["lists"]["A"]]
    def extra(stage):
        for index, item in enumerate(data.get("extra_steps", [])):
            if item["before"] == stage:
                steps.append(_step(f"extra_{index}", item["prompt"], "instruction", stage=stage))
    def study(name, trial):
        for index, item in enumerate(data["lists"][name], 1):
            prompt = data["prompts"]["study"].format(word=item["word"])
            instruction = data.get("study_instructions", {}).get(f"A{trial}" if name == "A" else "B")
            if index == 1 and instruction:
                prompt = instruction + "\n" + prompt
            steps.append(_step(f"{name}{trial}_study_{index:02d}", prompt + "\nReply only Ready.",
                               "ready", "Ready", phase="study", list=name, trial=trial, serial_position=index, material=item["word"]))
    def recall(sid, prompt, words, phase, **meta):
        followup = data.get("recall_followup_prompt")
        if followup:
            meta["followup_id"] = sid + "_followup"
        steps.append(_step(sid, prompt + "\nReturn only a JSON array of the words recalled, in recall order. Use [] if none.",
                           "cvlt_recall", words, phase=phase, all_a_words=words_a, **meta))
        if followup:
            steps.append(_step(sid + "_followup", followup + "\nReturn only a JSON array of any additional words recalled, in recall order. Use [] if none.",
                               "cvlt_recall_followup", words, phase=phase, all_a_words=words_a,
                               recall_id=sid, **{key: value for key, value in meta.items() if key != "followup_id"}))
    for trial in range(1, 6):
        extra(f"A{trial}"); study("A", trial)
        recall(f"A{trial}_recall", data["prompts"]["recall_a"], words_a, "learning_recall", list="A", trial=trial)
    extra("B"); study("B", 1)
    recall("B_recall", data["prompts"]["recall_b"], [x["word"] for x in data["lists"]["B"]], "interference_recall", list="B")
    for interval in ("short", "long"):
        extra(interval)
        if interval == "long":
            for index, item in enumerate(data.get("gap_steps", {}).get("before_long", [])):
                steps.append(_step(f"long_gap_{index}", item["prompt"], "instruction", phase="gap", delay_equivalence=False))
        recall(f"{interval}_free", data["prompts"][f"{interval}_free"], words_a, f"{interval}_free", list="A")
        for index, category in enumerate(data["cue_order"], 1):
            words = [x["word"] for x in data["lists"]["A"] if x["category"] == category]
            recall(f"{interval}_cued_{index}", data["prompts"]["cued"].format(category=category), words,
                   f"{interval}_cued", list="A", category=category)
    extra("recognition")
    for index, probe in enumerate(data["recognition"], 1):
        prompt = data["prompts"]["recognition"].format(word=probe["word"])
        if index == 1 and data.get("recognition_instructions"):
            prompt = data["recognition_instructions"] + "\n" + prompt
        steps.append(_step(f"recognition_{index:02d}", prompt + "\nReply only YES or NO.",
                           "cvlt_recognition", probe["is_old"], phase="recognition", word=probe["word"], is_old=probe["is_old"],
                           source_item_type=probe.get("record_form_item_type")))
    if data.get("forced_choice"):
        extra("forced_choice")
        for index, item in enumerate(data.get("gap_steps", {}).get("before_forced_choice", [])):
            steps.append(_step(f"forced_gap_{index}", item["prompt"], "instruction", phase="gap", delay_equivalence=False))
        for index, item in enumerate(data["forced_choice"], 1):
            steps.append(_step(f"forced_{index:02d}", f"Which word appeared in List A?\n1. {item['choices'][0]}\n2. {item['choices'][1]}\nReply only 1 or 2.",
                               "cvlt_forced", item["answer"], phase="forced_choice", cnphuman_endpoint_available=False))
    return _result([_episode("cvlt_standard", "cvlt_word_list", "standard_full_history", steps)], provenance)


def _read_drm(path):
    import pyreadstat  # existing project .deps/public_human_structure_20261015
    frame, metadata = pyreadstat.read_sav(path)
    study = frame[(frame.RunningBlock == "study") & (frame.sen.str.strip() != "")]
    probes = frame[frame.RunningBlock == "test"]
    if len(study) != 64 or len(probes) != 80:
        raise ValueError(f"Unexpected original DRM material counts: {len(study)}/{len(probes)}")
    if Counter(probes.oldnew) != Counter({"old": 40, "CL": 8, "newunrelated": 8, "newassoc": 20, "newtheme": 4}):
        raise ValueError("Original DRM probe classes differ from 40/8/8/20/4")
    return study.to_dict("records"), probes.to_dict("records")


DRM_TEST_INSTRUCTIONS = (
    "Decide whether each test word appeared in the study items, replying OLD or NEW. "
    "After OLD, classify the judgement as FAMILIAR (it feels familiar without a specific detail), "
    "REMEMBER (you recall study details), or GUESS. After NEW, classify it as "
    "UNFAMILIAR (it does not feel familiar), RECALL_TO_REJECT (you recall study information "
    "that rules it out), or WOULD_REMEMBER (you believe this item would have stood out if studied). "
    "For example, recalling that you studied plane rather than train can support RECALL_TO_REJECT; "
    "believing you would recognize your own name if it had appeared can support WOULD_REMEMBER."
)


def _load_drm_city(config):
    """Accept a verified word set for new randomization or an existing code mapping.

    A source citation is recorded for review; reading JSON does not itself verify
    that citation. Neither candidate upstream pool is silently installed here.
    """
    base = Path(config["materials_dir"])
    if config.get("drm_city_words"):
        path = Path(config["drm_city_words"])
    elif config.get("drm_city_mapping"):
        path = Path(config["drm_city_mapping"])
    else:
        word_path = base / "drm_city_words.json"
        path = word_path if word_path.is_file() else base / "drm_city_mapping.json"
    if not path.is_file():
        return None
    data = json.loads(path.read_text())
    if not isinstance(data.get("source_reference"), str) or not data["source_reference"].strip():
        raise ValueError("DRM CITY requires a genuine original-source reference for the eight-word set")
    if "words" in data:
        words = data["words"]
        if not isinstance(words, list):
            raise ValueError("DRM CITY words must be a list of eight distinct original words")
        mode = "verified_word_set_new_ids"
        mapping = None
    else:
        mapping = data.get("mapping", {})
        if not isinstance(mapping, dict) or set(mapping) != {str(x) for x in range(1, 9)}:
            raise ValueError("Legacy DRM CITY mapping requires eight code-to-word entries")
        words = list(mapping.values())
        mode = "original_code_mapping"
    if (len(words) != 8 or any(not isinstance(x, str) or not x.strip() or x.strip().isdigit() for x in words)
            or len({_norm(x) for x in words}) != 8):
        raise ValueError("DRM CITY requires eight distinct nonnumeric words")
    return {"file": str(path), "source_reference": data["source_reference"], "mode": mode,
            "words": sorted(words, key=_norm), "mapping": mapping}


def build_drm(config):
    raw = Path(config["raw_root"]) / "ball2022_memory/stimuli"
    requested_variant = config.get("drm_variant", "words_and_sentences")
    if requested_variant not in ("sentence", "words_and_sentences"):
        raise ValueError(f"Unsupported DRM variant: {requested_variant}")
    sentence_only = requested_variant == "sentence"
    provenance = {"cohort": "Ball et al. 2022", "source": "https://osf.io/r94uy/",
                  "raw_directory": str(raw), "drm_variant": requested_variant,
                  "material_scope": "sentence; warning and no warning" if sentence_only else "word and sentence; warning and no warning",
                  "complete_planned_steps_per_model": 448 if sentence_only else 896,
                  "order": "one seeded list-blocked order, shared between warnings, not participant-level replay",
                  "instructions": "Text interface faithfully paraphrases paper pp.5–6; E-Prime instruction screens not available.",
                  "limitations": ["Full history remains accessible; human two-minute instruction delay is not reproduced by sleep.",
                                  "Follow-up labels are reported judgements, not evidence for human-like subjective recollection."]}
    episodes, blockers = [], []
    # An approved sentence-only preparation must never consume a CITY candidate.
    # The earlier combined protocol and its source blocker remain separately named.
    city = None if sentence_only else _load_drm_city(config)
    if sentence_only:
        provenance["sentence_material_file"] = str(raw / "Ball et al. (2021) - DRM Sentences.sav")
        provenance["human_endpoint_scope"] = {
            "variant": "sentences", "conditions": ["no_warning", "warning"],
            "author_delta_dprime_field": "DRMs",
            "human_rates_prefix": "DRMs_",
            "aggregation": "Each warning condition separately; no word/sentence mean or warning-condition pooling.",
        }
        provenance["deferred_word_variant"] = {
            "status": "blocked_pending_source_verification",
            "reason": "Ball word-version CITY eight-word study set is not source-verified; retained for later verification, not replaced by sentence words.",
            "material_file": str(raw / "Ball et al. (2021) - DRM Words.sav"),
            "conditions": ["words_no_warning", "words_warning"],
            "planned_steps_per_model": 448,
            "city_material_loaded": False,
        }
    if city:
        provenance["city_material"] = {k: city[k] for k in ("file", "source_reference", "mode")}
        if city["mode"] == "original_code_mapping":
            provenance["city_mapping"] = {"file": city["file"], "source_reference": city["source_reference"]}
    for variant in (("sentences",) if sentence_only else ("words", "sentences")):
        path = raw / f"Ball et al. (2021) - DRM {variant.title()}.sav"
        if not path.is_file():
            blockers.append(f"Missing original DRM {variant} SAV: {path}"); continue
        study, probes = _read_drm(path)
        if variant == "words" and city is None:
            blockers.append("DRM word CITY eight-word study set is not source-verified. The old 1–8 mapping is not required for new seeded ordering; no guessed words or automatic CITY/word-condition exclusion.")
            # Sentence episodes are useful previews, never silently a complete ready DRM.
            continue
        if variant == "words":
            city_rows = sorted((r for r in study if r["sentheme"] == "CITY"), key=lambda r: r["sennum"])
            if len(city_rows) != 8:
                raise ValueError("Original DRM word CITY placeholders must contain eight study rows")
            for index, row in enumerate(city_rows):
                if city["mode"] == "original_code_mapping":
                    row["sen"] = city["mapping"][str(int(row["CITY"]))]
                else:
                    row["sen"] = city["words"][index]
                    row["prepared_item_id"] = f"city_word_{index + 1:02d}"
        rng = random.Random(int(config.get("seed", 20261017)) + (0 if variant == "words" else 1))
        groups = defaultdict(list)
        for row in study:
            groups[row["sentheme"]].append(row)
        themes = sorted(groups); rng.shuffle(themes)
        ordered = []
        for theme in themes:
            items = sorted(groups[theme], key=lambda x: x["sennum"]); rng.shuffle(items); ordered.extend(items)
        probes = sorted(probes, key=lambda x: x["wordnum"]); rng.shuffle(probes)
        for warned in (False, True):
            condition = f"{variant}_{'warning' if warned else 'no_warning'}"
            instruction = "Study each item; your memory for its words will be tested later. Rate its meaning from 1 (very meaningless) to 7 (very meaningful)."
            if warned:
                instruction += (" Each block of eight items is associated with one common word that is not itself studied. "
                                "For example, bed, rest, and tired may suggest sleep. Try to identify the common word while studying. "
                                "Such related but unstudied words will appear at test; avoid calling them OLD.")
            steps = []
            for index, row in enumerate(ordered):
                prompt = (instruction + "\n" if index == 0 else "")
                if index % 8 == 0:
                    prompt += f"Study block {index // 8 + 1}.\n"
                prompt += f"{row['sen']}\nMeaning rating (1–7 only):"
                item_id = row.get("prepared_item_id")
                steps.append(_step(f"study_{index+1:02d}", prompt, "drm_rating", phase="study", material=row["sen"],
                                   theme=row["sentheme"], source_sennum=None if item_id else int(row["sennum"]),
                                   prepared_item_id=item_id, variant=variant))
            for index, row in enumerate(probes, 1):
                prompt = (DRM_TEST_INSTRUCTIONS + "\n" if index == 1 else "") + f"Test word: {row['word']}\nReply only OLD or NEW."
                meta = {"phase": "recognition", "word": row["word"], "probe_class": row["oldnew"],
                        "source_wordnum": int(row["wordnum"]), "primary56": row["oldnew"] in ("old", "CL", "newunrelated"),
                        "variant": variant, "warned": warned}
                steps.append(_step(f"recognition_{index:02d}", prompt, "drm_recognition", bool(row["correct"]), **meta))
                steps.append(_step(f"report_{index:02d}", "Classify your immediately preceding OLD/NEW judgement using the previously defined labels.",
                                   "drm_report", phase="report", recognition_id=f"recognition_{index:02d}", **{k:v for k,v in meta.items() if k != "phase"}))
            episodes.append(_episode(f"drm_{condition}", "drm_false_memory", condition, steps))
    return _result(episodes, provenance, blockers)


def _eerland_materials(raw):
    import openpyxl
    from docx import Document
    # DOCX has explicit author word-count separators, preserving its four sections.
    names = {"Indirect": "indirect", "Direct – agent last": "direct_agent_last", "Direct – agent first": "direct_agent_first"}
    stories, current, chunks = {}, None, []
    for paragraph in Document(raw / "Stories and word count.docx").paragraphs:
        text = paragraph.text.strip()
        if text in names:
            current = names[text]; stories[current] = []; chunks = []
        elif current and re.match(r"^\d+ words(?:,|$)", text):
            stories[current].append(" ".join(chunks)); chunks = []
        elif current and text:
            chunks.append(text)
    if len(stories) != 3 or any(len(x) != 4 or not all(x) for x in stories.values()):
        raise ValueError("Eerland DOCX must provide three complete four-section stories")
    workbook = openpyxl.load_workbook(raw / "Stimuli_design.xlsx", read_only=True, data_only=True)
    rows = list(workbook["Stimuli"].values)
    speakers = list(rows[0][5:9]); questions = []
    correction = "Who asked the others about their new students?"
    for rownum, row in enumerate(rows[1:], 2):
        if not row[3]:
            continue
        marked = [speakers[i] for i, mark in enumerate(row[5:9]) if str(mark).lower() == "x"]
        evidence = "Stimuli_design.xlsx answer x"
        if row[3] == correction and not marked:
            if not all("Olivia" in str(x) for x in row[:3]):
                raise ValueError("Eerland missing-key correction no longer matches three source versions")
            marked = ["Olivia"]
            evidence = "Three story versions at design row20 and author Data_AnalysisExp1 Q53 correct code3 (Olivia); original cell remains unchanged"
        if len(marked) != 1:
            raise ValueError(f"Eerland question at row {rownum} lacks a unique source-backed key")
        questions.append({"id": f"q_row{rownum}", "question": row[3], "theme": int(row[4]),
                          "answer": marked[0], "key_evidence": evidence, "source_row": rownum})
    workbook.close()
    if len(questions) != 48 or sorted(Counter(q["theme"] for q in questions).values()) != [4] * 12:
        raise ValueError("Eerland must have 48 questions in twelve four-item theme groups")
    return stories, speakers, questions


def build_source_monitoring(config):
    raw = Path(config["raw_root"]) / "eerland2018_source_memory"
    provenance = {"cohort": "Eerland & Zwaan 2018 Exp1/2", "source": "https://osf.io/rk548/", "raw_directory": str(raw),
                  "stories_file": "Stories and word count.docx", "questions_file": "Stimuli_design.xlsx",
                  "scope": "48 source questions in two complementary 24-item forms per each of three stories",
                  "order": "seeded two-of-four per theme, independent sessions, shared forms/order across stories",
                  "limitations": ["New cohort does not replace Ball/CNP same-participant correlations.",
                                  "Full earlier story remains in model context; humans could not look back.",
                                  "Complete question-bank coverage is not every human random order or every two-of-four combination."]}
    missing = [str(raw / f) for f in ("Stories and word count.docx", "Stimuli_design.xlsx") if not (raw / f).is_file()]
    if missing:
        return _result([], provenance, [f"Missing original Eerland source: {x}" for x in missing])
    stories, speakers, questions = _eerland_materials(raw)
    rng = random.Random(int(config.get("seed", 20261017)))
    forms = {"A": [], "B": []}
    for theme in sorted({q["theme"] for q in questions}):
        group = [q for q in questions if q["theme"] == theme]; rng.shuffle(group)
        forms["A"].extend(group[:2]); forms["B"].extend(group[2:])
    for form in forms.values():
        rng.shuffle(form)
    episodes = []
    for variant, paragraphs in stories.items():
        for form_name, form in forms.items():
            condition = f"{variant}_form{form_name}"
            steps = [_step(f"read_{i}", ("Read the story carefully; questions follow after all four sections.\n" if i == 1 else "") +
                           f"Story section {i}/4:\n{text}\nReply only Ready.", "ready", "Ready", phase="study", section=i, material=text)
                     for i, text in enumerate(paragraphs, 1)]
            for q in form:
                options = "\n".join(f"{i}. {name}" for i, name in enumerate(speakers, 1))
                steps.append(_step(q["id"], q["question"] + "\n" + options + "\nReply only the speaker's name or option number.",
                                   "source_choice", q["answer"], phase="test", theme=q["theme"], source_row=q["source_row"],
                                   options=speakers, key_evidence=q["key_evidence"], story_variant=variant, form=form_name))
            episodes.append(_episode(f"sm_{condition}", "source_monitoring", condition, steps))
    provenance["story_words"] = {key: sum(len(p.split()) for p in value) for key, value in stories.items()}
    provenance["answer_key_correction"] = next(q for q in questions if q["source_row"] == 20)
    return _result(episodes, provenance)


def build(config):
    return {"cvlt_word_list": build_cvlt(config), "drm_false_memory": build_drm(config), "source_monitoring": build_source_monitoring(config)}


def messages(episode, step, history):
    if episode.get('presentation'):
        if episode['presentation'] != 'compact_history_v1':
            raise ValueError(f"Unsupported memory presentation: {episode['presentation']}")
        if episode['task'] == 'drm_false_memory':
            from .compact_drm import messages as compact_messages
        elif episode['task'] == 'cvlt_word_list':
            from .compact_cvlt import messages as compact_messages
        else:
            raise ValueError(f"Unsupported compact memory task: {episode['task']}")
        return compact_messages(episode, step, history)
    result = [{"role": "system", "content": episode["system"]}]
    for index, record in enumerate(history):
        prior = record["step"]
        result.append({"role": "user", "content": _prompt_for(prior, history[:index])})
        result.append({"role": "assistant", "content": record["response"]})
    result.append({"role": "user", "content": _prompt_for(step, history)})
    return result


def _prompt_for(step, history):
    if step["kind"] != "drm_report":
        return step["prompt"]
    rid = step["meta"]["recognition_id"]
    record = next((r for r in reversed(history) if r["step"]["id"] == rid), None)
    choice = record["score"].get("parsed") if record and record["score"].get("valid") else None
    if choice is True:
        return "You answered OLD. Classify that judgement: FAMILIAR, REMEMBER, or GUESS. Reply with one label only."
    if choice is False:
        return "You answered NEW. Classify that judgement: UNFAMILIAR, RECALL_TO_REJECT, or WOULD_REMEMBER. Reply with one label only."
    return "The preceding OLD/NEW response could not be classified. Reply only UNCLASSIFIABLE."


def _strict_score_step(step, response):
    kind, gold = step["kind"], step.get("gold")
    if kind in ("cvlt_recall", "cvlt_recall_followup"):
        try:
            parsed = json.loads(response)
        except (json.JSONDecodeError, TypeError):
            return {"parsed": None, "valid": False, "is_correct": None, "reason": "expected JSON word array"}
        if not isinstance(parsed, list) or any(not isinstance(x, str) or not x.strip() for x in parsed):
            return {"parsed": parsed, "valid": False, "is_correct": None, "reason": "expected JSON array of nonempty strings"}
        words = [_norm(x) for x in parsed]; target = {_norm(x) for x in gold}
        correct = set(words) & target
        intrusions = [x for x in words if x not in target]
        return {"parsed": words, "valid": True, "is_correct": len(correct) == len(target) and not intrusions,
                "correct_count": len(correct), "target_count": len(target), "proportion": len(correct) / len(target),
                "intrusions": intrusions, "repetitions": len(words) - len(set(words)),
                "out_of_category_a_words": [x for x in intrusions if x in {_norm(a) for a in step["meta"].get("all_a_words", [])}]}
    if kind in ("cvlt_recognition", "drm_recognition"):
        options = {"yes": True, "no": False} if kind == "cvlt_recognition" else {"old": True, "new": False}
    elif kind == "drm_rating":
        options = {str(i): i for i in range(1, 8)}
    elif kind == "drm_report":
        options = {x.lower(): x for x in ("FAMILIAR", "REMEMBER", "GUESS", "UNFAMILIAR", "RECALL_TO_REJECT", "WOULD_REMEMBER")}
    elif kind == "source_choice":
        options = {_norm(x): x for x in step["meta"]["options"]}
        options.update({str(i): x for i, x in enumerate(step["meta"]["options"], 1)})
    elif kind == "cvlt_forced":
        options = {"1": 1, "2": 2}
    elif kind == "ready":
        options = {"ready": "Ready"}
    elif kind == "instruction":
        return {"parsed": response, "valid": bool(str(response).strip()), "is_correct": None}
    else:
        raise ValueError(f"Unsupported memory step kind: {kind}")
    parsed = _choice(response, options); valid = parsed is not None
    return {"parsed": parsed, "valid": valid, "is_correct": (parsed == gold if valid and gold is not None else None)}


def score_step(step, response):
    """Parse the whole final answer; preserve the original parser for comparison."""
    original = _strict_score_step(step, response)
    text = _answer_text(response)
    kind = step["kind"]
    if kind in ("cvlt_recall", "cvlt_recall_followup"):
        # Only a complete, sole JSON code block is an array wrapper. Multiple
        # arrays, explanatory prose, missing fences and unquoted words stay invalid.
        match = re.fullmatch(r"```(?:json)?[ \t]*\n(.*?)\n```", str(response).strip(), re.S | re.I)
        if match:
            text = match.group(1)
    elif kind == "source_choice":
        options = {str(i): name for i, name in enumerate(step["meta"]["options"], 1)}
        label = _label_with_option(text, options)
        if label is not None:
            text = label
    elif kind == "drm_report":
        label = re.fullmatch(r'The label for this response is "([A-Z_]+)"[.!]?', text, re.I)
        if label:
            text = label.group(1)
    scored = _strict_score_step(step, text)
    if kind == "drm_report" and text.upper() == "UNCLASSIFIABLE":
        # This is the requested response to an unparseable preceding judgement,
        # not an OLD/NEW response or one of the six behavioural report categories.
        scored.update(parsed=None, valid=False, is_correct=None,
                      unclassifiable_response=True,
                      reason="Failure-branch acknowledgement; no behavioural report value")
    return {**scored, "strict_original_score": original}


def _cvlt_recall_trials(rows):
    """Combine an initial recall and its one permitted continuation in order."""
    lookup = {(r["episode_id"], r["step"]["id"]): r for r in rows}
    trials = []
    for row in rows:
        step = row["step"]
        if step["kind"] != "cvlt_recall":
            continue
        trial = {"episode_id": row["episode_id"], "step_id": step["id"],
                 "phase": step["meta"]["phase"]}
        followup_id = step["meta"].get("followup_id")
        if not followup_id:
            trials.append({**trial, **row["score"]})
            continue
        followup = lookup.get((row["episode_id"], followup_id))
        if followup and (followup["step"]["kind"] != "cvlt_recall_followup"
                         or followup["step"]["meta"].get("recall_id") != step["id"]):
            raise ValueError(f"CVLT recall continuation does not match its initial trial: {followup_id}")
        initial_score = row["score"]
        followup_score = followup["score"] if followup else None
        if initial_score["valid"] and followup_score and followup_score["valid"]:
            combined = score_step(step, json.dumps(initial_score["parsed"] + followup_score["parsed"]))
        else:
            combined = {"parsed": None, "valid": False, "is_correct": None,
                        "reason": "Initial recall or required continuation is missing/invalid; no trial total imputed."}
        trials.append({**trial, **combined, "followup_step_id": followup_id,
                       "initial_score": initial_score, "followup_score": followup_score,
                       "response_steps_observed": 1 + int(followup is not None),
                       "response_steps_expected": 2})
    return trials


def summarize(records):
    """Descriptive endpoints only. No human/model effect comparison or p-values.

    Invalid responses remain explicit, never silently scored as NEW or zero recall.
    Rates show valid-response denominator plus attempted/invalid counts.
    """
    groups = defaultdict(list)
    for record in records:
        groups[(record.get("model"), record["task"], record["condition"])].append(record)
    output = []
    for (model, task, condition), rows in groups.items():
        summary = {"model": model, "task": task, "condition": condition, "attempted_steps": len(rows),
                   "invalid_steps": sum(not r["score"]["valid"] for r in rows)}
        if task == "cvlt_word_list":
            recalls = _cvlt_recall_trials(rows)
            summary["recall_steps"] = recalls
            by_phase = defaultdict(list)
            for r in recalls:
                by_phase[r["phase"]].append(r)
            learning = by_phase["learning_recall"]
            summary["learning_total_correct_out_of_80"] = (
                sum(r["correct_count"] for r in learning)
                if len(learning) == 5 and all(r["valid"] for r in learning) else None)
            summary["recall_endpoints"] = {}
            for phase, expected, denominator in (("interference_recall", 1, 16), ("short_free", 1, 16),
                                                 ("long_free", 1, 16), ("short_cued", 4, 16), ("long_cued", 4, 16)):
                items = by_phase[phase]
                total = sum(r["correct_count"] for r in items) if len(items) == expected and all(r["valid"] for r in items) else None
                summary["recall_endpoints"][phase] = {"correct": total, "denominator": denominator,
                    "proportion": total / denominator if total is not None else None,
                    "steps_observed": len(items), "expected_steps": expected}
            summary["recognition"] = _recognition_summary([r for r in rows if r["step"]["kind"] == "cvlt_recognition"], "is_old")
            typed = [r for r in rows if r["step"]["kind"] == "cvlt_recognition" and r["step"]["meta"].get("source_item_type")]
            if typed:
                summary["recognition_by_source_type"] = _recognition_summary(typed, "source_item_type")
        elif task == "drm_false_memory":
            recognition = [r for r in rows if r["step"]["kind"] == "drm_recognition"]
            counts = _recognition_summary(recognition, "probe_class")
            summary["recognition"] = counts
            variants = {r["step"]["meta"]["variant"] for r in recognition if "variant" in r["step"]["meta"]}
            warnings = {r["step"]["meta"]["warned"] for r in recognition if "warned" in r["step"]["meta"]}
            if len(variants) > 1 or len(warnings) > 1:
                raise ValueError("DRM summaries require one stimulus variant and one warning condition; no pooled word/sentence endpoint.")
            if variants:
                variant = next(iter(variants))
                summary["variant"] = variant
                summary["warning"] = ("warning" if next(iter(warnings)) else "no_warning") if warnings else None
                summary.update(_drm_structure_endpoints(counts, variant))
            primary = ("old", "CL", "newunrelated")
            if all(str(k) in counts and counts[str(k)]["valid"] > 0 for k in primary):
                z = NormalDist().inv_cdf
                rates = {k: (counts[k]["old_responses"] + .5) / (counts[k]["valid"] + 1) for k in primary}
                summary["hautus_valid_response_denominator"] = {
                    "dprime_lure": z(rates["old"]) - z(rates["CL"]),
                    "dprime_new": z(rates["old"]) - z(rates["newunrelated"]),
                    "delta_dprime": z(rates["CL"]) - z(rates["newunrelated"]),
                    "all_primary_answers_valid": all(counts[k]["invalid"] == 0 for k in primary)}
            lookup = {(r["episode_id"], r["step"]["id"]): r for r in recognition}
            reports = Counter(); inconsistent = 0
            for r in rows:
                if r["step"]["kind"] != "drm_report": continue
                previous = lookup.get((r["episode_id"], r["step"]["meta"]["recognition_id"]))
                expected = {"FAMILIAR", "REMEMBER", "GUESS"} if previous and previous["score"].get("parsed") is True else {"UNFAMILIAR", "RECALL_TO_REJECT", "WOULD_REMEMBER"}
                if not previous or not previous["score"]["valid"] or not r["score"]["valid"] or r["score"]["parsed"] not in expected:
                    inconsistent += 1
                else: reports[r["score"]["parsed"]] += 1
            summary["followup_labels"] = dict(reports)
            summary["invalid_or_inconsistent_followups"] = inconsistent
        elif task == "source_monitoring":
            test = [r for r in rows if r["step"]["kind"] == "source_choice"]
            valid = [r for r in test if r["score"]["valid"]]
            summary.update({"questions": len(test), "valid_questions": len(valid), "invalid_questions": len(test)-len(valid),
                            "correct": sum(r["score"]["is_correct"] for r in valid),
                            "accuracy_valid": sum(r["score"]["is_correct"] for r in valid)/len(valid) if valid else None,
                            "confusion": {str(k):dict(v) for k,v in _source_confusion(valid).items()}})
        output.append(summary)
    return output


def _drm_structure_endpoints(counts, variant):
    """Source-aligned endpoints, published only for all 56 valid primary probes.

    Ball's DRMs is a difference in d-prime, not recognition accuracy. The native
    false-alarm contrast has the opposite orientation. The 24 additional probes
    stay in recognition summaries but cannot enter the human 56-probe endpoint.
    """
    prefix = {"sentences": "drm_sentence", "words": "drm_word"}[variant]
    expected = {"old": 40, "CL": 8, "newunrelated": 8}
    complete = all(k in counts and counts[k]["attempted"] == n and counts[k]["valid"] == n
                   for k, n in expected.items())
    endpoints = {f"{prefix}_{name}": None for name in
                 ("native", "accuracy_analyzed56", "balanced_accuracy", "delta_dprime")}
    if complete:
        hit, lure, new = (counts[k]["old_response_rate"] for k in expected)
        z = NormalDist().inv_cdf
        corrected = {k: (counts[k]["old_responses"] + .5) / (n + 1) for k, n in expected.items()}
        endpoints.update({
            f"{prefix}_native": new - lure,
            f"{prefix}_accuracy_analyzed56": (40 * hit + 8 * (1 - lure) + 8 * (1 - new)) / 56,
            f"{prefix}_balanced_accuracy": (hit + (1 - lure) + (1 - new)) / 3,
            f"{prefix}_delta_dprime": z(corrected["CL"]) - z(corrected["newunrelated"]),
        })
    return {"complete_primary56": complete, "primary56_expected_counts": expected,
            "structure_endpoints": endpoints,
            "structure_endpoint_scope": "One stimulus variant and warning condition; no word/sentence or warning-condition average. Invalid or incomplete primary probes yield null endpoints."}


def _recognition_summary(rows, field):
    counts = defaultdict(lambda: {"attempted": 0, "valid": 0, "invalid": 0, "old_responses": 0})
    for r in rows:
        c = counts[str(r["step"]["meta"][field])]; c["attempted"] += 1
        if r["score"]["valid"]:
            c["valid"] += 1; c["old_responses"] += int(r["score"]["parsed"])
        else: c["invalid"] += 1
    for c in counts.values(): c["old_response_rate"] = c["old_responses"] / c["valid"] if c["valid"] else None
    return dict(counts)


def _source_confusion(rows):
    confusion = defaultdict(Counter)
    for r in rows: confusion[r["step"]["gold"]][r["score"]["parsed"]] += 1
    return confusion
