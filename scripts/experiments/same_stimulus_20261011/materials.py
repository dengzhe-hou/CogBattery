"""Build the P9 same-stimulus materials (login-node CPU, no model call).

    cd <repository root>
    OMP_NUM_THREADS=1 MKL_NUM_THREADS=1 OPENBLAS_NUM_THREADS=1 nice -n 10 \
        python -m scripts.experiments.same_stimulus_20261011.materials

Outputs (results/same_stimulus_20261011/materials/):
- nback_episodes.jsonl          80 SRO adaptive n-back blocks (40 at load 2, 40 at load 3), letters as presented
- human_nback_trials.jsonl      the sampled participants' responses on exactly those blocks
- human_nback_matched.json      human endpoints on the sampled blocks + full-sample pooled reference
- strachan_items.jsonl          Strachan 2024 false-belief lists A and B (9 items each), text verbatim
- human_strachan_items.json     per-item human proportion correct (scores_human.csv) and per-subject totals
- strachan_scorer_validation.json  automated rule vs the authors' coding on their coded responses
- strachan_scorer_validation_r1.json  same check for scorer revision r1; written 2026-10-11 22:18 by the
                                     earlier interface `--scorer-only` (no revision argument) with the r1 code
- strachan_scorer_validation_r1a.json same check for scorer revision r1a (`--scorer-only r1a`)
- materials_summary.json        counts, sampling log, call counts
"""
import csv
import json
import os
import random
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from . import protocol as P
from . import scoring as SC

ROOT = Path(__file__).resolve().parents[3]
RAW = Path(os.environ.get('COGBATTERY_RAW', ROOT / 'data' / 'raw'))
OUT = Path(os.environ.get('COGBATTERY_SAME_STIMULUS_OUT', ROOT / 'results/same_stimulus_20261011/materials'))
SRO_SCRIPTS = ROOT / 'scripts/reanalysis/human_reference_20261009'
SRO_FULL_REF = ROOT / 'results/reanalysis/human_reference_20261009/sro/summary.json'
STRACHAN_XLSX = (RAW / 'strachan2024_tom/osf_fwj6v_20261011/Data/'
                     'Dataset - Theory of Mind in LLMs and Humans (Strachan et al 2023).xlsx')
STRACHAN_HUMAN = RAW / 'strachan2024_tom/osf_fwj6v_20261009/scored_data/scores_human.csv'
SEED = 20261011
K_PER_LOAD = {3: 40, 2: 40}   # load 3 sampled first (fewer eligible participants), then load 2
PILOT_PER_LOAD = 3
STRACHAN_STATE = {'A': ['T', 'T', 'F', 'F', 'T', 'F', 'FC', 'T', 'FC'],
                  'B': ['F', 'F', 'T', 'T', 'F', 'T', 'FC', 'F', 'TC']}
STRACHAN_SOURCE = ['old', 'old', 'old', 'old', 'new', 'new', 'catch', 'new', 'catch']
STRACHAN_PILOT = {'A': [1, 2, 3, 4, 5, 6]}  # list A positions 1-6 = 3 true-belief + 3 false-belief items
WORKBOOK_ONLY_SESSIONS = {('A', 4), ('A', 8)}
LOC_LABEL = {1: ('top drawer of the freezer', 'bottom drawer'), 2: ('in the sand next to the sandcastle', 'under the bucket'),
             3: ('next to the daffodils', 'next to the roses'), 4: ('in the playbox', 'under the sofa'),
             5: ('/shared_folder/tmp', '/shared_folder/photos'), 6: ('side pocket', 'front pocket'),
             8: ('top drawer of the cabinet', 'bottom drawer'), 9: ('tin in the living room', 'shelf in the kitchen')}


def now():
    return datetime.now(ZoneInfo('Asia/Tokyo')).isoformat(timespec='seconds')


def write_json(name, obj):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=1, allow_nan=False) + '\n')


def write_jsonl(name, rows):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in rows))


# ---------------------------------------------------------------- SRO n-back
def sro_blocks():
    """QC-passed adaptive blocks at loads 2 and 3, with the authors' QC (sro_human.qc)."""
    sys.path.insert(0, str(SRO_SCRIPTS))
    import sro_human as S  # noqa: E402  (asserts OMP/MKL/OPENBLAS == 1)
    d = S.screened('adaptive_n_back')
    qc = S.qc_report['adaptive_n_back']
    d = d[d.exp_stage == 'adaptive'].copy()
    blocks = {}
    n_checked = 0
    for (w, b), g in d.groupby(['worker_id', 'block_num'], sort=False):
        g = g.sort_values('time_elapsed', kind='stable')
        n = int(g.load.iloc[0])
        assert (g.load == n).all()
        if n not in (2, 3):
            continue
        seq = g.stim.tolist()
        for st, html in zip(seq, g.stimulus.tolist()):
            m = re.search(r'center-text">([A-Za-z])</div>', html)
            assert m and m.group(1) == st, (w, b, st, html)
        types = SC.classify(seq, n)
        for t, cr in zip(types, g.correct_response.tolist()):
            assert (cr == 37) == (t['kind'] == 'target'), (w, b)
            n_checked += 1
        assert len(seq) == 20 + n
        keys = g.key_press.tolist()
        codes = [None if k == -1 else (k == 37) for k in keys]
        assert all(k in (-1, 37, 40) for k in keys)
        blocks[(w, int(b))] = dict(worker_id=w, block_num=int(b), load=n, sequence=seq, types=types,
                                   human_codes=codes, human_rt=[float(x) for x in g.rt.tolist()],
                                   human_correct=[int(x) for x in g.correct.tolist()],
                                   n_lure=sum(t['kind'] == 'lure' for t in types))
    return blocks, qc, n_checked


def sample_blocks(blocks):
    rng = random.Random(SEED)
    used = set()
    picked = {}
    log = {}
    for load in (3, 2):
        elig = defaultdict(list)
        for (w, b), x in blocks.items():
            if x['load'] == load and x['n_lure'] >= 1:
                elig[w].append(b)
        n_blocks = sum(1 for x in blocks.values() if x['load'] == load)
        workers = sorted(w for w in elig if w not in used)
        chosen = rng.sample(workers, K_PER_LOAD[load])
        picked[load] = [(w, rng.choice(sorted(elig[w]))) for w in chosen]
        used.update(chosen)
        log[load] = dict(blocks_at_load=n_blocks, eligible_blocks=sum(len(v) for v in elig.values()),
                         eligible_participants=len(elig), candidates_after_excluding_used=len(workers),
                         sampled=K_PER_LOAD[load])
    return picked, log


def build_nback():
    blocks, qc, n_checked = sro_blocks()
    picked, log = sample_blocks(blocks)
    episodes, human = [], []
    for load in (2, 3):
        for k, (w, b) in enumerate(picked[load]):
            x = blocks[(w, b)]
            eid = f'sro_nb_L{load}_{w}_b{b:02d}'
            seq = x['sequence']
            episodes.append(dict(
                episode_id=eid, paradigm='n_back_sro', load=load, n_trials=len(seq), sequence=seq,
                trial_kinds=[t['kind'] for t in x['types']], lm1=[t['lm1'] for t in x['types']],
                lp1=[t['lp1'] for t in x['types']], overview=P.nback_overview(load, len(seq), seq[0]),
                source=dict(dataset='SRO adaptive_n_back', worker_id=w, block_num=b),
                sample_index=k, pilot=k < PILOT_PER_LOAD))
            human.append(dict(episode_id=eid, worker_id=w, block_num=b, load=load,
                              codes=x['human_codes'], rt=x['human_rt'], correct=x['human_correct']))
    # human endpoints on exactly these blocks (each block from a different participant)
    matched = {}
    for load in (2, 3):
        hs = [h for h in human if h['load'] == load]
        eps = [[dict(kind=k, code=c) for k, c in zip(e['trial_kinds'], h['codes'])]
               for e, h in zip([e for e in episodes if e['load'] == load], hs)]
        res = SC.endpoints(eps)
        per = []
        for e, h, tr in zip([e for e in episodes if e['load'] == load], hs, eps):
            r = SC.endpoints([tr])['R1']
            per.append((r['lure']['rate'], r['control']['rate']))
        cnt = Counter('gt' if (a is not None and c is not None and a > c + 1e-12) else
                      'lt' if (a is not None and c is not None and a < c - 1e-12) else
                      'tie' if (a is not None and c is not None) else 'undefined' for a, c in per)
        mb = defaultdict(lambda: [0, 0])
        for e, h in zip([e for e in episodes if e['load'] == load], hs):
            for m, v in SC.mback_consistency(e['sequence'], h['codes']).items():
                mb[m][0] += v['agree']
                mb[m][1] += v['n']
        res['participants_lureFA_vs_controlFA_R1'] = dict(cnt)
        res['mback_consistency_pooled'] = {str(m): dict(agree=a, n=n, rate=a / n if n else None) for m, (a, n) in mb.items()}
        matched[load] = res
    # full-sample pooled reference with the same function (all QC-passed blocks at the load)
    full = {}
    for load in (2, 3):
        eps = [[dict(kind=t['kind'], code=c) for t, c in zip(x['types'], x['human_codes'])]
               for x in blocks.values() if x['load'] == load]
        r = SC.endpoints(eps)
        full[load] = {k: r[k] for k in ('R1', 'R2', 'dprime_episode_mean', 'dprime_episodes_used', 'n_episodes',
                                        'n_trials', 'no_response_rate')}
        full[load]['n_participants'] = len({w for (w, b), x in blocks.items() if x['load'] == load})
    # The full-sample reference is a human-side summary. Material-only builds run without it.
    ref = (json.loads(SRO_FULL_REF.read_text())['n_back']['results']['authors_rule_missed_excluded']['by_load']
           if SRO_FULL_REF.is_file() else None)
    human_out = dict(
        note='Human endpoints on the sampled blocks (one block per participant) and two full-sample references. '
             'R1: rates over trials with a key press; R2: all trials, a missed response counts as not MATCH. '
             'd prime per block (log-linear, >=5 targets and >=5 non-targets with a response), averaged.',
        matched_blocks={str(k): v for k, v in matched.items()},
        full_sample_pooled_same_function={str(k): v for k, v in full.items()},
        full_sample_per_subject_mean_reference=(dict(path=str(SRO_FULL_REF), by_load={k: ref[k] for k in ('2', '3')})
                                               if ref else None),
        qc=qc, target_rule_checked_trials=n_checked)
    return episodes, human, human_out, log, qc


# ---------------------------------------------------------------- Strachan
def read_strachan():
    import openpyxl
    wb = openpyxl.load_workbook(STRACHAN_XLSX, read_only=False)
    sheets = {}
    for lst in 'AB':
        ws = wb[f'ToM-B False Belief - {lst}']
        heads = [ws.cell(1, 3 + 2 * k).value for k in range(9)]
        assert all(h for h in heads) and ws.cell(2, 3).value == 'Answer1' and ws.cell(2, 4).value == 'Score1'
        rows = []
        for r in range(3, ws.max_row + 1):
            who = ws.cell(r, 1).value
            if who is None:
                continue
            rows.append(dict(who=who, session=ws.cell(r, 2).value,
                             answers=[ws.cell(r, 3 + 2 * k).value for k in range(9)],
                             scores=[ws.cell(r, 4 + 2 * k).value for k in range(9)]))
        sheets[lst] = dict(heads=heads, rows=rows)
    return sheets


def read_human_scores():
    out = {}
    with STRACHAN_HUMAN.open() as f:
        for row in csv.DictReader(f):
            if not row['task'].startswith('False Belief'):
                continue
            lst = row['task'][-1]
            vals = [row[f'score{i}'] for i in range(1, 52)]
            out[(lst, int(row['item']))] = dict(state=row['trial_state'], source=row['source'],
                                                scores=[None if v == 'NA' else int(float(v)) for v in vals])
    return out


def build_strachan():
    sheets = read_strachan()
    hs = read_human_scores()
    items, human_items = [], []
    for lst in 'AB':
        heads = sheets[lst]['heads']
        preambles = {h.split('\n\n', 1)[0] for h in heads}
        assert len(preambles) == 1
        hrows = [r for r in sheets[lst]['rows'] if r['who'] == 'Human']
        # alignment check: workbook human session s of list A is scores_human.csv column score{s};
        # list B session s is score{25+s}.  Workbook A sessions 4 and 8 are coded in the workbook
        # but NA in the csv (the paper's N=49); they stay out of the human rates.
        off = 0 if lst == 'A' else 25
        for r in hrows:
            col = int(r['session']) - 1 + off
            wb_scores = [None if v is None else int(v) for v in r['scores']]
            csv_scores = [hs[(lst, p)]['scores'][col] for p in range(1, 10)]
            if (lst, int(r['session'])) in WORKBOOK_ONLY_SESSIONS:
                assert all(v is None for v in csv_scores), (lst, r['session'])
            else:
                assert wb_scores == csv_scores, (lst, r['session'], wb_scores, csv_scores)
        for k in range(9):
            pos = k + 1
            h = hs[(lst, pos)]
            assert h['state'] == STRACHAN_STATE[lst][k], (lst, pos, h['state'])
            src = 'catch' if h['state'] in ('FC', 'TC') else h['source']
            assert src == STRACHAN_SOURCE[k], (lst, pos, h['source'])
            sc = [s for s in h['scores'] if s is not None]
            text = heads[k].strip()
            items.append(dict(item_id=f'{lst}{pos}', list=lst, position=pos, story=pos, state=h['state'],
                              source=src, main_analysis=h['state'] in ('T', 'F'), text=text,
                              correct_location=('not surprised (ice melts)' if pos == 7 else
                                                LOC_LABEL[pos][0] if SC.correct_is_orig(pos, lst) else LOC_LABEL[pos][1]),
                              other_location=(None if pos == 7 else
                                              LOC_LABEL[pos][1] if SC.correct_is_orig(pos, lst) else LOC_LABEL[pos][0]),
                              pilot=pos in STRACHAN_PILOT.get(lst, [])))
            human_items.append(dict(item_id=f'{lst}{pos}', list=lst, position=pos, state=h['state'],
                                    n=len(sc), correct=sum(sc), rate=sum(sc) / len(sc)))
    # per-subject totals (each participant did one list)
    subj = []
    for lst in 'AB':
        cols = range(51)
        for c in cols:
            fb = [hs[(lst, p)]['scores'][c] for p in range(1, 10) if hs[(lst, p)]['state'] == 'F']
            tb = [hs[(lst, p)]['scores'][c] for p in range(1, 10) if hs[(lst, p)]['state'] == 'T']
            if any(v is None for v in fb + tb):
                assert all(v is None for v in fb + tb)
                continue
            subj.append(dict(list=lst, column=f'score{c + 1}', fb_correct=sum(fb), fb_n=len(fb), tb_correct=sum(tb), tb_n=len(tb)))
    tot = {}
    for st in ('F', 'T'):
        xs = [h for h in human_items if h['state'] == st]
        tot[st] = dict(correct=sum(h['correct'] for h in xs), n=sum(h['n'] for h in xs))
        tot[st]['rate'] = tot[st]['correct'] / tot[st]['n']
    human = dict(source=str(STRACHAN_HUMAN), licence='CC BY-NC 4.0, research-internal use only',
                 items=human_items, subjects=subj, n_subjects=dict(Counter(s['list'] for s in subj)),
                 totals_FB_TB=tot,
                 subjects_at_ceiling=dict(fb=sum(s['fb_correct'] == s['fb_n'] for s in subj),
                                          tb=sum(s['tb_correct'] == s['tb_n'] for s in subj), n=len(subj)))
    return items, human, sheets


def validate_scorer(sheets):
    """Automated rule vs the authors' coding on every coded response in the two FB sheets."""
    tab = defaultdict(lambda: [0, 0])
    paths = Counter()
    dis = []
    for lst in 'AB':
        heads = sheets[lst]['heads']
        for r in sheets[lst]['rows']:
            for k in range(9):
                a, s = r['answers'][k], r['scores'][k]
                if a is None or s is None:
                    continue
                st = STRACHAN_STATE[lst][k]
                pred, path, _ = SC.score_fb(k + 1, lst, str(a), heads[k])
                grp = ('main_FB_TB' if st in 'TF' else 'catch', 'human' if r['who'] == 'Human' else 'llm')
                tab[grp][0] += int(pred == int(s))
                tab[grp][1] += 1
                paths[(path, pred == int(s))] += 1
                if pred != int(s):
                    dis.append(dict(item=f'{lst}{k + 1}', state=st, who=r['who'], session=r['session'],
                                    authors_code=int(s), rule=pred, path=path, answer=str(a)))
    return dict(
        note='Rule developed on these same coded responses (r0 before any model call; r1/r1a revised after pilot 25552, before any formal call). In-sample agreement, not a held-out test.',
        agreement={f'{g[0]}|{g[1]}': dict(agree=v[0], n=v[1], rate=v[0] / v[1]) for g, v in sorted(tab.items())},
        by_path={f'{p}|{"agree" if ok else "disagree"}': n for (p, ok), n in sorted(paths.items())},
        disagreements=dis)


def main():
    for k in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
        assert os.environ.get(k) == '1', f'{k} must be 1 on the login node'
    t0, c0 = time.time(), time.process_time()
    episodes, human_trials, human_nb, log, qc = build_nback()
    items, human_fb, sheets = build_strachan()
    val = validate_scorer(sheets)
    write_jsonl('nback_episodes.jsonl', episodes)
    write_jsonl('human_nback_trials.jsonl', human_trials)
    write_json('human_nback_matched.json', human_nb)
    write_jsonl('strachan_items.jsonl', items)
    write_json('human_strachan_items.json', human_fb)
    write_json('strachan_scorer_validation.json', val)
    nb_calls = {str(L): sum(e['n_trials'] for e in episodes if e['load'] == L) for L in (2, 3)}
    nb_pilot = {str(L): sum(e['n_trials'] for e in episodes if e['load'] == L and e['pilot']) for L in (2, 3)}
    kinds = {str(L): dict(Counter(k for e in episodes if e['load'] == L for k in e['trial_kinds'])) for L in (2, 3)}
    summary = dict(
        created_at=now(), seed=SEED, sro_qc=qc, sampling=dict(rule=(
            'Population: SRO adaptive n-back participants passing the authors\' QC (sro_human.qc). '
            'Eligible block: adaptive stage, load 2 or 3, at least one lure trial. Load 3 first: sample '
            f'{K_PER_LOAD[3]} participants (random.Random({SEED}).sample over sorted worker ids), one eligible '
            f'block each (rng.choice over sorted block numbers); then load 2: {K_PER_LOAD[2]} participants '
            'from those not already used, one block each. 80 distinct participants. Pilot = first '
            f'{PILOT_PER_LOAD} sampled blocks per load.'), by_load={str(k): v for k, v in log.items()}),
        nback=dict(episodes=len(episodes), by_load={str(L): sum(e['load'] == L for e in episodes) for L in (2, 3)},
                   trials_by_load=nb_calls, trial_kinds_by_load=kinds,
                   pilot_episodes=sum(e['pilot'] for e in episodes), pilot_trials_by_load=nb_pilot),
        strachan=dict(items=len(items), main_analysis_items=sum(i['main_analysis'] for i in items),
                      pilot_items=[i['item_id'] for i in items if i['pilot']]),
        calls_per_model=dict(nback=sum(nb_calls.values()), strachan_session=len(items), strachan_isolated=len(items),
                             total=sum(nb_calls.values()) + 2 * len(items)),
        pilot_calls_per_model=dict(nback=sum(nb_pilot.values()), strachan_session=sum(i['pilot'] for i in items),
                                   strachan_isolated=sum(i['pilot'] for i in items),
                                   total=sum(nb_pilot.values()) + 2 * sum(i['pilot'] for i in items)),
        scorer_agreement=val['agreement'],
        resource=dict(wall_seconds=round(time.time() - t0, 2), cpu_seconds=round(time.process_time() - c0, 2)))
    write_json('materials_summary.json', summary)
    print(json.dumps({k: summary[k] for k in ('nback', 'strachan', 'calls_per_model', 'pilot_calls_per_model',
                                              'scorer_agreement', 'resource')}, ensure_ascii=False, indent=1))


def main_scorer_only(revision):
    """Recompute agreement of the current scoring code on the authors' coding without rewriting any
    material file.  `revision` is only the file label; it does not select scoring code, so pass the
    name of the revision now in scoring.py.  Used as `--scorer-only r1a` (2026-10-11).  The r1 file was
    written earlier by `--scorer-only` without an argument; the r0 and r1 files stay as written and an
    existing file is not overwritten."""
    for k in ('OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'OPENBLAS_NUM_THREADS'):
        assert os.environ.get(k) == '1', f'{k} must be 1 on the login node'
    name = f'strachan_scorer_validation_{revision}.json'
    if (OUT / name).exists():
        raise SystemExit(f'{OUT / name} exists; validation files of earlier revisions are kept')
    val = validate_scorer(read_strachan())
    val['revision'] = revision
    val['created_at'] = now()
    write_json(name, val)
    print(json.dumps(dict(agreement=val['agreement'], by_path=val['by_path']), indent=1))


if __name__ == '__main__':
    if sys.argv[1:2] == ['--scorer-only']:
        main_scorer_only(sys.argv[2])
    else:
        main()
